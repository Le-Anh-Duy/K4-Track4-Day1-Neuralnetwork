"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

FORMULA_COLS = ("step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise")
TEMPLATE_ROWS = 60  # mẫu có công thức cho dòng 2..61
LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _clean(v):
    """Đổi NaN/inf -> None (Excel không có NaN) và tuple -> list (JSON)."""
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if isinstance(v, tuple):
        return list(v)
    return v


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json. Trả về đường dẫn."""
    out = Path(results_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{result['cfg']['exp_id']}.json"
    payload = {k: result[k] for k in ("cfg", "history", "summary")}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=1, ensure_ascii=False)  # NaN của lần chạy phân kỳ được giữ (Python đọc lại được)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir (sắp theo exp_id)."""
    results = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p, encoding="utf-8") as f:
            r = json.load(f)
        if {"cfg", "history", "summary"} <= r.keys():
            r["cfg"]["hidden"] = tuple(r["cfg"]["hidden"])
            results.append(r)
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả -> một dòng của bảng (khoá trùng tên cột). Chỉ truyền eval_scores cho baseline và cấu hình cuối."""
    cfg, s = result["cfg"], result["summary"]
    extra = []
    if cfg.get("scheduler"):
        extra.append(f"scheduler={cfg['scheduler']}")
    if cfg["optimizer"] == "sgd_momentum" and cfg.get("momentum", 0.9) != 0.9:
        extra.append(f"momentum={cfg['momentum']}")
    if cfg["optimizer"] in ("adam", "adamw"):
        extra.append("betas=(0.9,0.999), eps=1e-8")
    if cfg["clip_norm"] is not None and "clip_frac" in result["history"]:
        cf = result["history"]["clip_frac"]
        extra.append(f"clip_frac TB={sum(cf) / len(cf):.2f}")
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=LOSS_NAMES[cfg["loss"]], optimizer=OPT_NAMES[cfg["optimizer"]], lr=cfg["lr"],
        weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(map(str, cfg["hidden"])), dropout=cfg["dropout"],
        clip_norm="none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        **{k: s.get(k) for k in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
                                 "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB", "diverged")},
        eval_acc=eval_scores["accuracy"] if eval_scores else None,
        eval_macro_f1=eval_scores["macro_f1"] if eval_scores else None,
        figure_file=f"figures/{cfg['exp_id']}.png",
        notes="; ".join([n for n in [notes, *extra] if n]),
    )
    return {k: _clean(v) for k, v in row.items()}


def write_xlsx(rows: list[dict], template_path: str, out_path: str, seed_ids: list[str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu (từ dòng 2), ghi exp_id baseline vào sheet "Seeds"
    (cột A, dòng 2..6), rồi lưu thành out_path. Giữ nguyên công thức; mở bằng Excel/LibreOffice để tính lại.
    """
    import openpyxl

    if len(rows) > TEMPLATE_ROWS:
        raise ValueError(f"mẫu chỉ có công thức cho {TEMPLATE_ROWS} dòng, đang có {len(rows)}")
    ids = [r["exp_id"] for r in rows]
    assert len(set(ids)) == len(ids), "exp_id bị trùng"

    wb = openpyxl.load_workbook(template_path)  # KHÔNG data_only=True (sẽ mất công thức)
    ws = wb["Experiments"]
    col_of = {c.value: c.column for c in ws[1] if c.value}
    data_cols = [k for k in col_of if k not in FORMULA_COLS]
    for i in range(TEMPLATE_ROWS):  # xoá dữ liệu mẫu (dòng baseline minh hoạ) rồi ghi lại
        for k in data_cols:
            ws.cell(row=2 + i, column=col_of[k]).value = rows[i].get(k) if i < len(rows) else None

    if seed_ids is not None:
        ws_s = wb["Seeds"]
        for i in range(5):  # dòng 2..6
            ws_s.cell(row=2 + i, column=1).value = seed_ids[i] if i < len(seed_ids) else None
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
